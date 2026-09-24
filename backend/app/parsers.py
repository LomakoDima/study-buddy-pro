from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from fastapi import HTTPException
from pptx import Presentation


SUPPORTED_EXTENSIONS = {".pdf": "PDF", ".docx": "DOCX", ".pptx": "PPTX"}


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    unit_count: int
    file_type: str


def _clean(value: str) -> str:
    value = value.replace("\x00", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    extension = Path(filename).suffix.lower()
    file_type = SUPPORTED_EXTENSIONS.get(extension)
    if not file_type:
        raise HTTPException(status_code=415, detail="Only PDF, DOCX and PPTX files are supported")
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty")

    try:
        if extension == ".pdf":
            with pymupdf.open(stream=content, filetype="pdf") as document:
                units = [
                    f"[Page {index + 1}]\n{page.get_text('text')}"
                    for index, page in enumerate(document)
                ]
        elif extension == ".docx":
            document = Document(io.BytesIO(content))
            units = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
                units.append("\n".join(rows))
        else:
            presentation = Presentation(io.BytesIO(content))
            units = []
            for index, slide in enumerate(presentation.slides):
                texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
                units.append(f"[Slide {index + 1}]\n" + "\n".join(texts))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="The document is damaged or could not be read") from exc

    cleaned_units = [_clean(unit) for unit in units if _clean(unit)]
    text = "\n\n".join(cleaned_units)
    if len(text) < 20:
        raise HTTPException(
            status_code=422,
            detail="No readable text was found. Scanned/image-only documents are not supported yet.",
        )
    return ParsedDocument(text=text, unit_count=len(cleaned_units), file_type=file_type)


def logical_chunks(text: str, max_chars: int) -> list[str]:
    blocks = [block.strip() for block in text.split("\n\n") if block.strip()]
    chunks: list[str] = []
    current: list[str] = []
    current_length = 0

    for block in blocks:
        if len(block) > max_chars:
            paragraphs = [block[i : i + max_chars] for i in range(0, len(block), max_chars)]
        else:
            paragraphs = [block]
        for paragraph in paragraphs:
            extra = len(paragraph) + (2 if current else 0)
            if current and current_length + extra > max_chars:
                chunks.append("\n\n".join(current))
                current = []
                current_length = 0
            current.append(paragraph)
            current_length += len(paragraph) + (2 if len(current) > 1 else 0)
    if current:
        chunks.append("\n\n".join(current))
    return chunks
