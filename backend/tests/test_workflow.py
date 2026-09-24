import asyncio
import io

from docx import Document
from fastapi.testclient import TestClient
import pymupdf
from pptx import Presentation

from app.auth import AuthUser
from app.db import Database
from app.parsers import logical_chunks, parse_document
from app.services import SummaryService


class TestSummarizer:
    configured = True

    async def summarize(self, filename, text, mode, on_progress=None):
        assert filename == "lecture.docx"
        assert "retrieval practice" in text
        assert mode == "notes"
        if on_progress:
            on_progress(70)
        return {
            "title": "Memory and Retrieval",
            "tags": ["Memory", "Learning"],
            "key_points": ["Retrieval practice improves long-term recall."],
            "terms": ["Retrieval practice"],
            "sections": [{"heading": "Overview", "text": "Active recall strengthens memory."}],
            "qa": [],
        }


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Learning Science", level=1)
    document.add_paragraph("Spacing and retrieval practice improve long-term learning and recall.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_upload_to_summary_library_and_delete(tmp_path):
    database = Database(tmp_path / "test.db")
    service = SummaryService(database, TestSummarizer())
    content = make_docx()
    parsed = parse_document("lecture.docx", content)
    user = AuthUser(42, "Ava", "ava")

    document = service.create_document(user, "lecture.docx", len(content), parsed)
    summary = service.create_summary(user, document["id"], "notes")
    asyncio.run(service.generate_summary(summary["id"]))

    ready = service.get_summary(user.id, summary["id"])
    assert ready["status"] == "ready"
    assert ready["payload"]["title"] == "Memory and Retrieval"
    assert service.list_summaries(user.id)[0]["id"] == summary["id"]

    service.delete_summary(user.id, summary["id"])
    assert service.list_summaries(user.id) == []


def test_logical_chunks_keep_all_text():
    source = "First section.\n\n" + ("A" * 30) + "\n\nLast section."
    chunks = logical_chunks(source, 20)
    assert len(chunks) >= 3
    assert "First section." in "\n\n".join(chunks)
    assert "Last section." in "\n\n".join(chunks)


def test_pdf_and_pptx_parsers():
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Cell membranes regulate transport and signaling.")
    parsed_pdf = parse_document("cells.pdf", pdf.tobytes())
    assert parsed_pdf.file_type == "PDF"
    assert "Cell membranes" in parsed_pdf.text

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Supply and demand"
    slide.placeholders[1].text = "Equilibrium occurs where quantities are equal."
    stream = io.BytesIO()
    presentation.save(stream)
    parsed_pptx = parse_document("economics.pptx", stream.getvalue())
    assert parsed_pptx.file_type == "PPTX"
    assert "Equilibrium" in parsed_pptx.text


def test_http_upload_to_summary_workflow(tmp_path):
    from app import main

    database = Database(tmp_path / "api.db")
    summarizer = TestSummarizer()
    main.database = database
    main.summarizer = summarizer
    main.service = SummaryService(database, summarizer)
    user = AuthUser(77, "Local Student", "local")
    main.app.dependency_overrides[main.current_user] = lambda: user

    try:
        with TestClient(main.app) as client:
            upload = client.post(
                "/api/documents",
                files={
                    "file": (
                        "lecture.docx",
                        make_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            assert upload.status_code == 201

            create = client.post(
                f"/api/documents/{upload.json()['id']}/summaries",
                json={"mode": "notes"},
            )
            assert create.status_code == 202
            summary_id = create.json()["id"]

            stored = client.get(f"/api/summaries/{summary_id}")
            assert stored.status_code == 200
            assert stored.json()["status"] == "ready"
            assert client.get("/api/summaries").json()[0]["id"] == summary_id
    finally:
        main.app.dependency_overrides.clear()
